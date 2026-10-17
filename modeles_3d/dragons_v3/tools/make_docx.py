"""Assemble le docx final : Dragon_Pets_Scripts_Blender_v3.docx
Usage : bl/bin/python tools/make_docx.py  (depuis /home/claude/work)"""
import json, os, re
from PIL import Image
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

W = "/home/claude/work"
OUT = os.path.join(W, "Dragon_Pets_Scripts_Blender_v3.docx")
D = json.load(open(f"{W}/data/dragons.json"))
KIT = open(f"{W}/kit/kit.py", encoding="utf-8").read().rstrip() + "\n"
STAGES = [("BEBE", "Bébé", 4), ("ADO", "Ado", 7), ("ADULTE", "Adulte", 11)]
BUDGET = {"Commun": (3000, 4500, 6000), "Rare": (4000, 6000, 8000), "Épique": (5000, 8500, 13000),
          "Légendaire": (6000, 10000, 15000), "Mythique": (7000, 11000, 17000), "Secret": (7000, 11000, 17000),
          "Exclusif": (7000, 11000, 17000), "Mythe": (8000, 13000, 20000)}
DARK = RGBColor(0x26, 0x26, 0x26)
GREY = RGBColor(0x59, 0x59, 0x59)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2)
sec.top_margin = sec.bottom_margin = Cm(1.8)

st = doc.styles["Normal"]
st.font.name = "Calibri"; st.font.size = Pt(10.5); st.font.color.rgb = DARK
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(4)
for name, size, before in (("Title", 26, 0), ("Heading 1", 18, 18), ("Heading 2", 13, 12), ("Heading 3", 11, 8)):
    s = doc.styles[name]
    s.font.name = "Calibri"; s.font.size = Pt(size); s.font.bold = name != "Title"
    s.font.color.rgb = DARK
    rf = s.element.rPr.find(qn("w:rFonts"))
    if rf is not None:
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]
        rf.set(qn("w:ascii"), "Calibri"); rf.set(qn("w:hAnsi"), "Calibri")
    s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(4)

def shade(par, fill="F2F2F2"):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    pPr.append(shd)

def mono(text, size=7.5, fill="F4F4F4"):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text.rstrip("\n"))
    r.font.name = "Consolas"; r.font.size = Pt(size)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    shade(p, fill)
    return p

def para(text, size=None, color=None, italic=False, bold=False, after=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    r.italic = italic; r.bold = bold
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p

def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if " : " in it and len(it.split(" : ", 1)[0]) < 40:
            k, v = it.split(" : ", 1)
            p.add_run(k).bold = True
            p.add_run(" : " + v)
        else:
            p.add_run(it)
        p.paragraph_format.space_after = Pt(2)

def table(rows, widths, size=9):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.width = Cm(widths[j]); c.text = ""
            r = c.paragraphs[0].add_run(str(val)); r.font.size = Pt(size)
            if i == 0:
                r.bold = True
                tcPr = c._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "EDEDED"); tcPr.append(shd)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def field(run, code, placeholder=""):
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = code
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = placeholder
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    for x in (b, it, s, t, e): run._r.append(x)

def preview(slug, stage):
    a, b = f"{W}/renders/{slug}_{stage}.png", f"{W}/renders/{slug}_{stage}_side.png"
    if not (os.path.exists(a) and os.path.exists(b)):
        return None
    ims = [Image.open(x).convert("RGB") for x in (a, b)]
    h = 520
    ims = [i.resize((int(i.width * h / i.height), h)) for i in ims]
    can = Image.new("RGB", (sum(i.width for i in ims) + 10, h), "white")
    x = 0
    for i in ims:
        can.paste(i, (x, 0)); x += i.width + 10
    os.makedirs(f"{W}/docx_img", exist_ok=True)
    out = f"{W}/docx_img/{slug}_{stage}.jpg"
    can.save(out, quality=80)
    return out

def report(slug):
    try:
        return json.load(open(f"{W}/out/{slug}/report.json", encoding="utf-8"))["stages"]
    except Exception:
        return {}

def fix_prompt(t):
    for a, b in (("sur le Adulte", "sur l'Adulte"), ("sur le Ado ", "sur l'Ado "), (" le Adulte", " l'Adulte"),
                 (" le Ado ", " l'Ado "), (" du Adulte", " de l'Adulte"), (" du Ado ", " de l'Ado ")):
        t = t.replace(a, b)
    return re.sub(r"^```\w*\s*$", "", t.strip(), flags=re.M).strip()

# ------------------------------------------------------------------ page de titre
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run(); fr.font.size = Pt(9); fr.font.color.rgb = GREY; field(fr, "PAGE", "1")

p = doc.add_paragraph(style="Title"); p.add_run("Dragon Pets Simulator")
para("Scripts Blender des 42 dragons, version 3", size=15, color=GREY, after=2)
para("Un prompt et un script Python par dragon et par stade (Bébé, Ado, Adulte) : 126 modèles, "
     "plus détaillés que la version 2 et stylés selon leur rareté.", size=11, after=14)
para("Raboul · d'après le GDD v2.3 et l'exemple « Adulte Mirage » · octobre 2026", size=10, color=GREY, after=18)
doc.add_heading("Sommaire", level=2)
field(doc.add_paragraph().add_run(), 'TOC \\o "1-1" \\h \\z \\u',
      "Clic droit puis « Mettre à jour les champs » pour afficher le sommaire.")
doc.add_page_break()

# ------------------------------------------------------------------ mode d'emploi
doc.add_heading("Comment s'en servir", level=1)
doc.add_heading("Lancer un script dans Blender", level=3)
bullets([
    "Ouvre Blender, onglet Scripting, « Nouveau », colle tout le bloc de code du dragon voulu (de « import bpy » "
    "jusqu'à la dernière ligne), puis « Exécuter le script » (Alt+P).",
    "Le dragon arrive dans sa propre collection (par exemple POUSSE_ADULTE). Si tu relances le script, seule cette "
    "collection est remplacée : les autres stades et les autres dragons de la scène ne bougent pas.",
    "La console affiche à la fin le nombre de triangles et le compare au budget du stade.",
])
doc.add_heading("Ce que fait chaque script", level=3)
bullets([
    "Même méthode que ton Adulte Mirage : blocs biseautés, éclats plats pour les crinières et les crêtes, cristaux "
    "hexagonaux, plaques de ventre séparées, bracelets et griffes, corps longs sur une courbe Catmull-Rom.",
    "Hiérarchie prête pour le futur squelette : un empty <TAG>_MASTER_ROOT avec HEAD_ROOT, BODY_ROOT, LEGS_ROOT, "
    "TAIL_ROOT, plus WINGS_ROOT (ailes ou nageoires) et HEAD2_ROOT / HEAD3_ROOT pour les hydres. Les pièces de la "
    "tête sont en coordonnées locales de HEAD_ROOT.",
    "Axes : avant = -Y, haut = +Z, unité = mètre. Hauteur finale : Bébé 4 m, Ado 7 m, Adulte 11 m, posé sur Z = 0.",
    "Pas de bpy.ops dans les boucles : les meshes sont créés avec bmesh / from_pydata. Testé dans Blender 5.2, "
    "compatible 4.x.",
    "Seed fixe : chaque exécution donne exactement le même dragon.",
])
doc.add_heading("Réglages en haut de chaque script", level=3)
bullets([
    "TARGET_HEIGHT : hauteur finale du dragon.",
    "BEVEL_SEGMENTS : 1 par défaut ; 0 enlève les biseaux (environ deux fois moins de triangles).",
    "ADD_FLOATING_FX : True ou False pour les éléments flottants (cristaux, braises, bulles...).",
    "Bloc de forme : chaque script a ses propres réglages (taille de tête, envergure, longueur de queue...). Les plus "
    "utiles sont rappelés sous chaque script, dans « Réglages utiles ».",
])
doc.add_heading("À quoi sert le prompt", level=3)
para("Chaque prompt suit la structure de ton prompt « Adulte Mirage » et décrit ce que fait le script en dessous : "
     "pièces jointes à donner, étape de validation avant le code, ce que le modèle doit reproduire, contraintes "
     "techniques et livraison. Tu peux le donner à une IA avec l'image de concept pour régénérer le dragon, ou le "
     "coller avec le script pour demander une retouche précise.")
para("Les images de concept n'existaient que pour les mondes 1 et 2. Pour les autres dragons, le prompt précise que "
     "la description fait foi si tu n'as pas d'image.", italic=True, color=GREY)

doc.add_heading("Ce qui change avec la rareté", level=1)
para("Plus un dragon est rare, plus il a de détails et d'effets. L'Adulte porte tout, l'Ado environ 60 %, le Bébé "
     "une touche.")
table([
    ["Rareté", "Ce que le modèle ajoute"],
    ["Commun", "Silhouette propre et lisible, 3 à 5 couleurs, pas d'émission, un signe distinctif fort."],
    ["Rare", "Motif secondaire (bandes, taches, veines), cornes travaillées, crête double, un accent lumineux discret."],
    ["Épique", "Silhouette spectaculaire (crinière en éclats, grande crête), éléments émissifs, petits ornements dorés, éléments flottants."],
    ["Légendaire", "Tout l'Épique + armure en métal (couronne, épaulières, plastron à gemme), aura derrière la tête, double rangée de pics, ornement de queue."],
    ["Mythique", "Matériau précieux sur une grande partie du corps, gemmes multiples, halo flottant, éclats en orbite."],
    ["Secret", "Identité étrange : parties semi-transparentes, fissures lumineuses, silhouette déchirée, yeux lumineux, fragments flottants."],
    ["Mythe", "Le plus spectaculaire : couronne, anneaux runiques flottants, halo cosmique, constellations lumineuses."],
    ["Exclusif", "Thème très marqué (arc-en-ciel pour Prisma, galaxie pour Nébula), motifs lumineux, effets à thème."],
], [3, 14])
doc.add_heading("Budget de triangles", level=3)
para("Compté après modificateurs, avec BEVEL_SEGMENTS = 1 et les éléments flottants activés. Pour comparer, ton "
     "exemple Mirage v3 faisait 27 900 triangles avec ses biseaux à 2 segments.")
table([["Rareté", "Bébé", "Ado", "Adulte"]] + [[k, *[f"{v:,}".replace(",", " ") for v in b]] for k, b in BUDGET.items()],
      [5, 4, 4, 4])

doc.add_heading("Liste des dragons", level=1)
rows = [["Monde", "Dragon", "Élément", "Rareté", "Morphologie", "Triangles B / A / Ad"]]
for d in D:
    rp = report(d["slug"])
    rows.append([d["monde"], d["nom"], d["element"], d["rarete"], d["morphologie"],
                 " / ".join(str(rp.get(s, {}).get("tris", "?")) for s, _, _ in STAGES)])
table(rows, [1.6, 2.6, 2.2, 2.4, 2.6, 4.4], size=8.5)

# ------------------------------------------------------------------ dragons
missing = []
for d in D:
    slug = d["slug"]
    doc.add_page_break()
    doc.add_heading(d["nom"], level=1)
    monde = "Exclusif (payant)" if d["monde"] == "Exclusif" else f"Monde {d['monde']}"
    para(f"{monde} · {d['element']} · {d['rarete']} · {d['archetype']} · {d['morphologie']}", color=GREY, after=2)
    if d.get("signes"):
        para("Signes distinctifs : " + d["signes"], size=9.5, after=6)
    rp = report(slug)
    for key, label, h in STAGES:
        src = f"{W}/out/{slug}/{key}.py"
        if not os.path.exists(src):
            missing.append(f"{slug}/{key}"); continue
        if key != "BEBE":
            doc.add_page_break()
        doc.add_heading(f"{label} {d['nom']}", level=2)
        info = rp.get(key, {})
        if info.get("tris"):
            para(f"{info['tris']} triangles · hauteur {h} m", size=9, color=GREY, after=2)
        img = preview(slug, key)
        if img:
            doc.add_picture(img, width=Cm(15.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            para("Vue de trois quarts et vue de profil, rendues dans Blender à partir du script ci-dessous (la lueur "
                 "des parties lumineuses n'apparaît pas sur ces aperçus).", size=8, color=GREY, italic=True)
        else:
            missing.append(f"{slug}/{key}_render")
        if info.get("notes"):
            para(info["notes"], size=9.5)
        pf = f"{W}/out/{slug}/{key}_prompt.md"
        doc.add_heading("Prompt", level=3)
        if os.path.exists(pf):
            mono(fix_prompt(open(pf, encoding="utf-8").read()), size=8, fill="FAFAFA")
        else:
            missing.append(f"{slug}/{key}_prompt")
        fname = f"{slug.lower()}_{label.lower().replace('é', 'e')}.py"
        doc.add_heading(f"Script Python Blender · {fname}", level=3)
        mono(open(src, encoding="utf-8").read().replace("# @@KIT@@", KIT, 1))
        if info.get("reglages"):
            doc.add_heading("Réglages utiles", level=3)
            bullets(info["reglages"])

doc.save(OUT)
print("OK", OUT, "manquants:", missing)
